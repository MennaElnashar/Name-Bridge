"""Independent labelling check. A person who did not write the policy answers, for sampled pairs,
"same person / different people / cannot tell" from the shown fields only. We then compare their
answers with the policy and with the author's gold labels.

  python -m src.labelling make            # writes labelling/ (sheet for the reviewer, answer key)
  python -m src.labelling score FILE      # FILE: the reviewer's completed .xlsx or .csv

The reviewer sheet shows name, email, student ID and organization only. It hides the policy's answer,
the author's gold label, the ground-truth entity id and the free-text note. Pair order and A/B side
are shuffled with a fixed seed so the sheet is reproducible."""
import json
import random
import sys
from pathlib import Path

import pandas as pd

from .pipeline import find_candidates, read_csv_preserving
from .policy import PolicyConfig

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "labelling"
SEED = 16
QUOTA = {"likely_match": 8, "likely_different": 6, "needs_review": 10}  # 24 pairs, about 10 minutes
SHOWN = ["name", "email", "student_id", "org"]
ANSWERS = {"Same person": "likely_match", "Different people": "likely_different", "Cannot tell": "needs_review"}
MISSING = "(not recorded)"
ALIASES = {"same person": "Same person", "same": "Same person", "different people": "Different people",
           "different": "Different people", "diff": "Different people", "cannot tell": "Cannot tell",
           "can't tell": "Cannot tell", "cant tell": "Cannot tell", "unsure": "Cannot tell", "unknown": "Cannot tell"}


def _gold(gold_file):
    g = read_csv_preserving(ROOT / "data" / gold_file)
    return {tuple(sorted((r.record_id_a, r.record_id_b), key=int)): r.gold for r in g.itertuples()}


def build():
    rng = random.Random(SEED)
    pool = []
    for data, gold_file in (("synthetic_names.csv", "gold_pairs.csv"), ("hard_cases.csv", "hard_gold_pairs.csv")):
        df = read_csv_preserving(ROOT / "data" / data)
        recs = {r["record_id"]: r for r in df.to_dict("records")}
        gold = _gold(gold_file)
        for p in find_candidates(df, cfg=PolicyConfig()).to_dict("records"):
            a, b = p["record_id_a"], p["record_id_b"]
            pool.append({"set": data, "a": a, "b": b, "ra": recs[a], "rb": recs[b], "policy": p["recommendation"],
                         "gold": gold.get(tuple(sorted((a, b), key=int)), ""),
                         "truth_same": recs[a]["entity_id"] == recs[b]["entity_id"]})
    chosen = []
    for rec, n in QUOTA.items():
        group = [x for x in pool if x["policy"] == rec]
        chosen += rng.sample(group, min(n, len(group)))
    rng.shuffle(chosen)
    sheet, key = [], []
    for i, x in enumerate(chosen, 1):
        swap = rng.random() < 0.5
        left, right = (x["rb"], x["ra"]) if swap else (x["ra"], x["rb"])
        row = {"Pair": f"P{i:02d}"}
        for side, r in (("A", left), ("B", right)):
            for f, label in zip(SHOWN, ["Name", "Email", "Student ID", "Organization"]):
                row[f"{side} {label}"] = r[f] if str(r[f]).strip() else MISSING
        row["Your answer"] = ""
        row["Comment (optional)"] = ""
        sheet.append(row)
        key.append({"pair": f"P{i:02d}", "set": x["set"], "record_id_a": x["a"], "record_id_b": x["b"],
                    "swapped": swap, "policy": x["policy"], "gold": x["gold"], "truth_same_person": x["truth_same"]})
    return pd.DataFrame(sheet), pd.DataFrame(key)


INSTRUCTIONS = [
    "NameBridge: independent labelling",
    "",
    "Thank you for helping. This takes about 10 minutes and checks my work against a second pair of eyes.",
    "",
    "What to do: for each pair on the 'Pairs' sheet, look at the two records and pick an answer in the 'Your answer' column:",
    "  Same person      - the information shown makes it likely these are one person",
    "  Different people - the information shown makes it likely these are two people",
    "  Cannot tell      - the information shown is not enough to decide either way",
    "",
    "Rules:",
    "  1. Use only what is on the sheet. Please do not look at the project, the app or its repository first.",
    "  2. Spellings that look alike are NOT proof on their own. Arabic names have many valid spellings,",
    "     and different people can share the same name.",
    "  3. Matching emails or student IDs are evidence; conflicting ones are evidence the other way.",
    "     '(not recorded)' means the field was empty, so it tells you nothing either way.",
    "  4. If you are unsure, 'Cannot tell' is a perfectly good answer. Do not guess to be helpful.",
    "  5. A comment is optional but useful when you were torn.",
    "",
    "All names and details are invented. Save the file and send it back. Thank you!",
]


def write_sheet(df, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Instructions"
    for i, line in enumerate(INSTRUCTIONS, 1):
        ws0.cell(row=i, column=1, value=line)
    ws0["A1"].font = Font(bold=True, size=14)
    ws0.column_dimensions["A"].width = 120
    ws = wb.create_sheet("Pairs")
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))
    head_fill = {"A": "DCE6F1", "B": "E2EFDA"}
    for c, name in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        fill = head_fill.get(name[:1]) if name[:2] in ("A ", "B ") else ("FFF2CC" if name == "Your answer" else "EEEEEE")
        cell.fill = PatternFill("solid", fgColor=fill)
    widths = {"Pair": 7, "Your answer": 20, "Comment (optional)": 40}
    for c, name in enumerate(df.columns, 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = widths.get(name, 24 if "Email" in name else 20)
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ans_col = list(df.columns).index("Your answer") + 1
    dv = DataValidation(type="list", formula1='"Same person,Different people,Cannot tell"', allow_blank=True)
    dv.error, dv.errorTitle = "Please choose one of the three answers.", "Answer"
    ws.add_data_validation(dv)
    col = ws.cell(row=1, column=ans_col).column_letter
    dv.add(f"{col}2:{col}{len(df) + 1}")
    ws.freeze_panes = "B2"
    wb.save(path)


def score(sheet_df, key_df):
    """Compare a reviewer's answers with the policy and the gold labels."""
    d = sheet_df[["Pair", "Your answer"]].merge(key_df, left_on="Pair", right_on="pair")
    d["Your answer"] = d["Your answer"].astype(str).str.strip().str.lower().map(ALIASES)  # accept natural wording
    d = d[d["Your answer"].notna()].copy()
    d["human"] = d["Your answer"].map(ANSWERS)
    d["gold"] = d["gold"].fillna("").astype(str)
    out = {"pairs_answered": int(len(d))}
    if d.empty:
        return out, d
    out["agree_with_policy"] = int((d.human == d.policy).sum())
    g = d[d.gold != ""]
    out["pairs_with_gold_label"] = int(len(g))
    out["agree_with_gold"] = int((g.human == g.gold).sum())
    out["human_said_same_policy_said_different"] = int(((d.human == "likely_match") & (d.policy == "likely_different")).sum())
    out["human_said_different_policy_said_match"] = int(((d.human == "likely_different") & (d.policy == "likely_match")).sum())
    out["policy_match_where_truth_is_different_people"] = int(((d.policy == "likely_match") & (d.truth_same_person.astype(str).str.lower() != "true")).sum())
    out["human_cannot_tell_where_policy_decided"] = int(((d.human == "needs_review") & (d.policy != "needs_review")).sum())
    out["human_decided_where_policy_needs_review"] = int(((d.human != "needs_review") & (d.policy == "needs_review")).sum())
    out["disagreements"] = d[d.human != d.policy][["Pair", "record_id_a", "record_id_b", "human", "policy", "gold"]].to_dict("records")
    return out, d


def _read_sheet(path):
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xls"):
        return pd.read_excel(p, sheet_name="Pairs", dtype=str).fillna("")
    return pd.read_csv(p, dtype=str, keep_default_na=False)


def main(argv):
    if argv[:1] == ["make"]:
        OUT.mkdir(exist_ok=True)
        sheet, key = build()
        write_sheet(sheet, OUT / "labelling_sheet.xlsx")
        sheet.to_csv(OUT / "labelling_sheet.csv", index=False, encoding="utf-8-sig")
        key.to_csv(OUT / "answer_key.csv", index=False)
        print(f"wrote {len(sheet)} pairs to {OUT}. Give the reviewer ONLY labelling_sheet.xlsx (or .csv).")
        return 0
    if argv[:1] == ["score"] and len(argv) == 2:
        key = pd.read_csv(OUT / "answer_key.csv", dtype=str, keep_default_na=False)
        res, _ = score(_read_sheet(argv[1]), key)
        (ROOT / "outputs" / "human_label_results.json").write_text(json.dumps(res, indent=2), encoding="utf-8")
        print(json.dumps(res, indent=2))
        return 0
    print(__doc__)
    return 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
